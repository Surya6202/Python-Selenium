import json
from configparser import ConfigParser
from datetime import datetime
# from time import sleep

import openpyxl
from selenium import webdriver
from selenium.webdriver import ActionChains
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait

chrome_options = webdriver.ChromeOptions()
chrome_options.add_experimental_option(name='detach', value=True)
chrome_options.add_argument('--start-maximized')
chrome_options.add_experimental_option(name="excludeSwitches", value=["enable-automation"])
chrome_options.add_argument('--disable-notifications')
# chrome_options.add_argument(argument='--headless')
file_path= r'D:\Programming\Python\PythonSelenium\downloads'
chrome_options.add_experimental_option('prefs', {'download.default_directory': file_path})

edge_options = webdriver.EdgeOptions()
edge_options.add_experimental_option(name='detach', value=True)
edge_options.add_argument(argument='--start-maximized')
edge_options.add_experimental_option(name="excludeSwitches", value=["enable-automation"])
edge_options.add_argument(argument='--disable-notifications')
# edge_options.add_argument(argument='--headless')
# edge_options.add_experimental_option('prefs', {'download.default_directory': file_path})

ad_block = ('const elements = document.getElementsByClassName("adsbygoogle adsbygoogle-noablate");'
            'while (elements.length > 0) elements[0].remove()')
# DateTime stamp:
date = datetime.now().strftime('%d-%m-%Y')
time = datetime.now().strftime('%H-%M-%S%f')[0:10]

webpage_shot = f'../screenshots/{date}_webpagescreenshot_{time}.jpeg'
webelement_shot = f'../screenshots/{date}_webelementscreenshot_{time}.jpeg'


def web_browser(browser: str) -> WebDriver:
    if browser.lower().__eq__('chrome'):
        return webdriver.Chrome(chrome_options)
    elif browser.lower().__eq__('edge'):
        return webdriver.Edge(edge_options)
    elif browser.lower().__eq__('firefox'):
        return webdriver.Firefox()
    else:
        raise Exception("Invalid browser")


def json_data(key: str):
    file = open(r'../files/test_data.json').read()
    return json.loads(file)[key]


def excel_data(sheet_name: str, row_num: int, column_num: int):
    return openpyxl.load_workbook(r'../files/test_data.xlsx')[sheet_name].cell(row_num, column_num).value


class SeleniumUtility:

    def __init__(self, driver_control: WebDriver):
        self.driver = driver_control

    def precondition(self, url: str, secs: int | float) -> None:
        self.driver.implicitly_wait(secs)
        self.driver.set_page_load_timeout(secs)
        self.driver.get(url=url)

    def element(self, locator: str, value: str) -> WebElement:
        return self.driver.find_element(by=locator, value=value)

    def elements(self, locator: str, value: str) -> list[WebElement]:
        return self.driver.find_elements(by=locator, value=value)

    def wait(self, secs: int | float, poll_period: int | float) -> WebDriverWait:
        return WebDriverWait(driver=self.driver, timeout=secs, poll_frequency=poll_period)

    def select(self, locator: str, value: str) -> Select:
        return Select(self.element(locator=locator, value=value))

    def actions(self) -> ActionChains:
        return ActionChains(self.driver)

    @staticmethod
    def config_data(section: str, option: str):
        file = ConfigParser()
        file.read(r'../files/test_data.ini')
        return file.get(section=section, option=option)

    @staticmethod
    def excel_data(sheet_name: str, row_num: int, column_num: int):
        return openpyxl.load_workbook(r'../files/test_data.xlsx')[sheet_name].cell(row_num, column_num).value

    @staticmethod
    def json_data(key: str):
        file = open(r'../files/test_data.json').read()
        return json.loads(file)[key]

    def postcondition(self) -> None:
        self.driver.quit()


# driver = web_browser('Chrome')
# s = SeleniumUtility(driver)
# s.precondition(s.json_data('url'), 30)
# s.element('id', 'user-name').send_keys(s.json_data('username'))
# s.element('id', 'password').send_keys(s.json_data('password'))
# s.element('id', 'login-button').click()
# sleep(2)
# s.postcondition()

# class Demo():
#     def __init__(self, driver_control: WebDriver):
#         self._driver = driver_control
#         self._username = 'user-name'
#         self._password = 'password'
#         self._button = 'login-button'
#
#     def enter_username(self, username: str):
#         self._driver.find_element(self._username).send_keys(username)
#
#     def enter_password(self, password: str):
#         self._driver.find_element(self._password).send_keys(password)
#
#     def click_button(self):
#         self._driver.find_element(self._button).click()
#
#
# driver = web_browser(json_data('browser'))
# s = SeleniumUtility(driver)
# s.precondition(s.json_data('url'), 30)
# d = Demo(driver_control=driver)
# d.enter_username(s.json_data('username'))
# d.enter_password(s.json_data('password'))
# d.click_button()

# sleep(5)
# s.postcondition()


# from seleniumpagefactory.Pagefactory import *
#
#
# class Demo(PageFactory):
#
#     def __init__(self, driver_control: WebDriver):
#         super().__init__()
#         self.driver = driver_control
#
#     locators = {
#         "username": ('ID', 'user-name'),
#         "password": ('ID', 'password'),
#         "button": ('ID', 'login-button')
#     }
#
#     def enter_username(self) -> PageFactory:
#         return self.username
#
#     def enter_password(self) -> PageFactory:
#         return self.password
#
#     def click_login(self) -> PageFactory:
#         return self.button
#
#
# def login():
#     for i in range(2, 5):
#         driver = web_browser(excel_data('Sheet1', 1, i))
#         s = SeleniumUtility(driver_control=driver)
#         s.precondition(s.json_data('url'), 30)
#         d = Demo(driver_control=driver)
#         d.enter_username().set_text(s.json_data('username'))
#         d.enter_password().set_text(s.json_data('password'))
#         d.click_login().click_button()
#         sleep(2)
#         s.postcondition()
#         sleep(2)
#
#
# login()
